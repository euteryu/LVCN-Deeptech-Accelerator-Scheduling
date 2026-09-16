from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "LVCN_Programme_Board_Runbook.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Aptos", size=None, color=None, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def add_bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color="243047")
    return paragraph


def add_step(doc, number, title, body, command=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    number_run = p.add_run(f"{number}. ")
    set_run_font(number_run, size=11, color="4F46E5", bold=True)
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, color="0F172A", bold=True)
    body_p = doc.add_paragraph(body)
    body_p.paragraph_format.left_indent = Inches(0.22)
    body_p.paragraph_format.space_after = Pt(4)
    for run in body_p.runs:
        set_run_font(run, size=10.5, color="334155")
    if command:
        code = doc.add_paragraph()
        code.paragraph_format.left_indent = Inches(0.22)
        code.paragraph_format.space_after = Pt(7)
        code.paragraph_format.line_spacing = 1.0
        run = code.add_run(command)
        set_run_font(run, name="Consolas", size=9, color="E2E8F0")
        p_pr = code._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "0F172A")
        p_pr.append(shd)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    tr_pr = hdr._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = header
        shade(cell, "1E1B4B")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, color="FFFFFF", bold=True)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cells[idx], "F5F7FB")
            set_cell_border(cells[idx])
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    set_run_font(run, size=9, color="243047")
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_run_font(run, size=15 if level == 1 else 11.5, color="0F172A", bold=True)
    return p


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("LVCN Programme Board Runbook")
    set_run_font(run, size=24, color="0F172A", bold=True)
    subtitle = doc.add_paragraph("Local testing, Supabase setup, deployment, and programme close-down")
    subtitle.paragraph_format.space_after = Pt(12)
    for run in subtitle.runs:
        set_run_font(run, size=11, color="64748B")

    intro = doc.add_paragraph("This is the current follow-up guide for the implemented LVCN Programme Board. It explains exactly how to see the UI locally first, how to test the privacy model, and what to configure when moving from demo mode to a real Supabase-backed workspace.")
    intro.paragraph_format.space_after = Pt(9)
    for run in intro.runs:
        set_run_font(run, size=10.5, color="334155")

    heading(doc, "What is already complete")
    for text in [
        "Responsive calendar with week, day, and agenda views; mobile starts in agenda view.",
        "Separate LVCN programme events, targeted opportunities, business meetings, company work, and private availability.",
        "Startup decisions save immediately; compulsory items use Acknowledged only.",
        "Admin pulse metrics, filters, named response lists, conflict warnings, create/edit/duplicate/cancel/delete, and CSV import review.",
        "Excel and print-ready PDF export, invite-only Supabase email-link auth, database RLS, migrations, seed scripts, and isolation checklist.",
    ]:
        add_bullet(doc, text)

    heading(doc, "Part 1  See and test the UI locally")
    add_step(doc, 1, "Open the project folder", "Open PowerShell and move into the folder that contains package.json.", "cd C:\\Users\\minse\\Downloads\\LVCN_WEBAPP")
    add_step(doc, 2, "Install the dependencies", "This is needed once, or again after package.json changes.", "npm install")
    add_step(doc, 3, "Create the local environment file", "The example file intentionally contains only public placeholders. Demo mode does not need a real Supabase project.", "Copy-Item .env.example .env.local")
    add_step(doc, 4, "Confirm demo mode", "Open .env.local and leave VITE_DEMO_MODE=true. Leave the Supabase URL and key blank for the first UI review.", "notepad .env.local")
    add_step(doc, 5, "Start the development server", "Keep this PowerShell window open. Vite will print the local address.", "npm run dev")
    add_step(doc, 6, "Open the site", "Visit the printed address, normally http://localhost:5173. The demo opens as the LVCN administrator, with representative cohort data already present.", "start http://localhost:5173")

    heading(doc, "What to click through first", 2)
    add_table(doc, ["Test", "Expected result"], [
        ("Admin calendar", "Programme pulse, five-day week grid, item-type colours, filters, and named response counts are visible."),
        ("Profile menu", "Switch to Jin Park / Seoul Labs to preview the startup experience. In real mode this switch is not available."),
        ("Startup privacy", "Seoul Labs sees its own targeted items and private busy blocks; other startups' names and busy indicators are absent."),
        ("Item drawer", "Select an item. Admin sees cohort responses and conflict summaries; startup sees only its own decision and private note."),
        ("Decision flow", "Choose Going, Interested, or Pass. The saved state and toast appear immediately."),
        ("Mobile", "Use browser DevTools device emulation at 390px wide. Navigation collapses, agenda is the default, and Add busy time remains available."),
        ("Exports", "Use Export → Excel workbook or Print-ready PDF. The export uses the currently visible filtered schedule."),
    ], widths=[1.55, 5.9])

    heading(doc, "Useful local commands", 2)
    add_table(doc, ["Command", "Purpose"], [
        ("npm run test", "Run the representative-data tests."),
        ("npm run lint", "Check code style and React hook rules."),
        ("npm run typecheck", "Run strict TypeScript checking."),
        ("npm run build", "Create the production bundle in dist/."),
        ("npm run preview", "Serve the built dist/ bundle locally for a production-like check."),
    ], widths=[1.55, 5.9])

    heading(doc, "Part 2  Move from demo mode to real Supabase")
    add_step(doc, 1, "Create a Supabase project", "Create a project and keep the Project URL and public anonymous key available. Do not use the service-role key in this app.")
    add_step(doc, 2, "Configure email-link redirects", "In Supabase Authentication → URL Configuration, add http://localhost:5173 and the eventual Cloudflare Pages URL to the redirect allow-list.")
    add_step(doc, 3, "Apply the schema and policies", "From the project folder, link the Supabase CLI project and push the migrations.", "supabase link --project-ref YOUR_PROJECT_REF\nsupabase db push\nsupabase db seed")
    add_step(doc, 4, "Replace example invite emails", "Edit supabase/seed.sql or run SQL in the Supabase editor. Every startup member gets exactly one organisation; LVCN admins use a null organisation_id.", "insert into public.allowed_invites (email, full_name, role, organisation_id)\nvalues ('founder@company.com', 'Founder Name', 'startup_member', 'ORGANISATION_UUID');")
    add_step(doc, 5, "Set production environment variables", "Set the real Supabase values and turn demo mode off. The anonymous key is safe to ship; RLS is the protection boundary.", "VITE_SUPABASE_URL=https://PROJECT_REF.supabase.co\nVITE_SUPABASE_ANON_KEY=PUBLIC_ANON_KEY\nVITE_DEMO_MODE=false")
    add_step(doc, 6, "Sign in once as an admin", "The database trigger provisions a profile only when the email is on allowed_invites. An unlisted email receives the explicit access-denied screen and no application data.")
    add_step(doc, 7, "Add representative database rows", "After the admin profile exists, run supabase/seed_demo.sql in the SQL editor. It adds sample events, conflict groups, responses, and private availability using a real created_by user.")

    heading(doc, "Part 3  Test data isolation")
    for text in [
        "Invite two different startup emails, and open the app in two separate private browser windows or browser profiles.",
        "As Startup A, create a targeted item or response note and add private busy time.",
        "As Startup B, confirm that Startup A's targeted item, response note, availability block, and busy indicator are not present.",
        "As an LVCN admin, confirm that both organisations' records and named conflicts are visible.",
        "Repeat the SQL checklist in supabase/tests/rls_isolation.sql after any RLS policy change. Never test with the postgres or service_role role because those bypass RLS.",
    ]:
        add_bullet(doc, text)

    heading(doc, "Part 4  Deploy to Cloudflare Pages")
    add_step(doc, 1, "Create the Pages project", "Connect the Git repository in Cloudflare Pages.")
    add_step(doc, 2, "Set the build settings", "Use npm run build as the build command and dist as the output directory.")
    add_step(doc, 3, "Add environment variables", "Add VITE_SUPABASE_URL, VITE_SUPABASE_ANON_KEY, and VITE_DEMO_MODE=false in the Pages project settings.")
    add_step(doc, 4, "Deploy and update redirects", "Deploy, then add the final Pages URL to Supabase Authentication → URL Configuration. public/_redirects provides the SPA fallback for direct links.")

    heading(doc, "Part 5  Programme close-down")
    for text in [
        "Export the final schedule and confirm who owns the exported files.",
        "Revoke sessions and remove active addresses from allowed_invites.",
        "Remove Cloudflare environment variables and delete or unpublish the Pages project.",
        "Retain or delete Supabase data according to LVCN policy. Deleting a Supabase project is irreversible, so confirm the archive first.",
    ]:
        add_bullet(doc, text)

    heading(doc, "If something looks wrong")
    add_table(doc, ["Symptom", "Check first"], [
        ("Blank page", "Stop and restart the Vite server, then hard-refresh. Run npm run typecheck to catch import errors."),
        ("Access denied in real mode", "Confirm the exact lowercase email is present in allowed_invites, then sign out and request a new magic link."),
        ("No real data", "Confirm VITE_DEMO_MODE=false, both Supabase variables are set, migrations are applied, and an admin profile exists."),
        ("Direct route 404 after deployment", "Confirm public/_redirects is included and the Pages output directory is dist."),
        ("Unexpected privacy result", "Test in separate browser sessions and inspect RLS with the documented checklist; never rely on frontend hiding."),
    ], widths=[1.85, 5.6])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("LVCN Programme Board  ·  Runbook")
    set_run_font(run, size=8, color="94A3B8")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
